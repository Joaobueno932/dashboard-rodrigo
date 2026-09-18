"""Bounded, read-only XLSX extraction; cached values only, with cell provenance."""

import math
import re
import unicodedata
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

from .catalog import CATALOG, Indicator

MAX_ROWS = 30000
MAX_COLS = 64
MONTHS = [
    "janeiro",
    "fevereiro",
    "marco",
    "abril",
    "maio",
    "junho",
    "julho",
    "agosto",
    "setembro",
    "outubro",
    "novembro",
    "dezembro",
]


class WorkbookError(ValueError):
    pass


def norm(value: Any) -> str:
    text = " ".join(str(value or "").split()).lower()
    return "".join(c for c in unicodedata.normalize("NFD", text) if unicodedata.category(c) != "Mn")


def number(value: Any) -> float | None:
    if value is None or (isinstance(value, str) and value.strip() in ("", "-", "—", "–")):
        return None
    if isinstance(value, bool):
        raise WorkbookError("Valor booleano em coluna numérica.")
    if isinstance(value, str):
        value = value.strip().replace("\u00a0", "")
        if "," in value:
            value = value.replace(".", "").replace(",", ".")
    try:
        result = float(value)
    except (TypeError, ValueError):
        raise WorkbookError("Valor não numérico ou fórmula sem resultado válido.") from None
    if not math.isfinite(result) or result < 0:
        raise WorkbookError("Valor negativo ou não finito em indicador não negativo.")
    return result


def validate_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 4000 or sum(e.file_size for e in entries) > 100 * 1024 * 1024:
                raise WorkbookError("Excel excede o limite de conteúdo descompactado (100 MB / 4.000 arquivos).")
            names = {e.filename for e in entries}
            if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
                raise WorkbookError("O arquivo não é uma planilha XLSX válida.")
            if any("vbaproject" in n.lower() or n.startswith("xl/externalLinks/") for n in names):
                raise WorkbookError("Macros e vínculos externos não são aceitos.")
            for entry in entries:
                if entry.flag_bits & 1:
                    raise WorkbookError("Planilhas criptografadas não são aceitas.")
                if entry.filename.endswith((".xml", ".rels")):
                    content = archive.read(entry)
                    if b"<!DOCTYPE" in content or b"<!ENTITY" in content:
                        raise WorkbookError("Estrutura XML não permitida.")
            if archive.testzip():
                raise WorkbookError("Excel corrompido.")
    except (zipfile.BadZipFile, OSError, RuntimeError):
        raise WorkbookError("Excel inválido ou corrompido.") from None


def find_table(rows: list[tuple[Any, ...]], spec: Indicator) -> tuple[int, dict[str, int], list[int]] | None:
    for index, row in enumerate(rows):
        headers = [norm(v) for v in row]
        if "ano" not in headers:
            continue
        values = [i for i, h in enumerate(headers) if re.search(spec.value, h)]
        if not values:
            continue
        dims: dict[str, int] = {h: headers.index(h) for h in ("ano", "mes", "casa", "unidade") if h in headers}
        if spec.kind != "ratio" and not spec.code.startswith("G.") and "casa" not in dims:
            continue
        if spec.code in ("A.1.1.1", "A.2.1.1", "A.2.2.1") and "mes" not in dims:
            continue
        if spec.category:
            cats = [i for i, h in enumerate(headers) if re.search(spec.category, h)]
            if len(cats) != 1:
                continue
            dims["category"] = cats[0]
        if spec.kind != "wide" and len(values) != 1:
            raise WorkbookError(f"{spec.code}: cabeçalhos de valores ambíguos.")
        return index, dims, values
    return None


def parse_workbook(path: Path) -> dict[str, Any]:
    validate_archive(path)
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except Exception:
        raise WorkbookError("Não foi possível ler o Excel.") from None
    try:
        if len(workbook.sheetnames) > 150:
            raise WorkbookError("Limite de 150 abas excedido.")
        # Formatting in the supplied workbook extends to XFD. Bound reads to the
        # documented supported data area, independently of misleading dimensions.
        sheets = []
        for sheet in workbook:
            preview = list(sheet.iter_rows(max_row=100, max_col=MAX_COLS, values_only=True))
            codes = {
                str(v).strip()
                for row in preview
                for v in row
                if isinstance(v, str) and re.fullmatch(r"[ASG]\.\d+\.\d+\.\d+", v.strip())
            }
            if codes.intersection(s.code for s in CATALOG):
                sheets.append((sheet, preview, codes))
        indicators: list[dict[str, Any]] = []
        warnings = []
        for spec in CATALOG:
            candidates = []
            for sheet, preview, codes in sheets:
                if spec.code in codes:
                    table = find_table(preview, spec)
                    if table:
                        candidates.append((sheet, preview, table))
            if len(candidates) != 1:
                raise WorkbookError(
                    f"{spec.code}: indicador ausente, estrutura incompatível ou tabelas ambíguas ({len(candidates)} correspondências)."
                )
            sheet, preview, (header, dims, columns) = candidates[0]
            if (sheet.max_row or 0) > MAX_ROWS:
                raise WorkbookError(f"{spec.code}: limite de {MAX_ROWS} linhas excedido.")
            meta = next(row for row in preview if any(str(v).strip() == spec.code for v in row))
            code_col = next(i for i, v in enumerate(meta) if str(v).strip() == spec.code)
            unit = str(meta[code_col + 3] or "")
            if not unit:
                raise WorkbookError(f"{spec.code}: unidade de medida não identificada.")
            # Table units are more specific than metadata for GEE (tCO2e).
            header_value = str(preview[header][columns[0]])
            match = re.search(r"\(([^()]+)\)\s*$", header_value)
            display_unit = match.group(1) if match else unit
            records: list[dict[str, Any]] = []
            duplicates: Counter[tuple[Any, ...]] = Counter()
            missing = 0
            has_month = "mes" in dims
            relevant = set(dims.values()) | set(columns)
            for row_no, row in enumerate(
                sheet.iter_rows(
                    min_row=header + 2,
                    max_row=min(sheet.max_row or MAX_ROWS, MAX_ROWS) + 1,
                    max_col=MAX_COLS,
                    values_only=True,
                ),
                header + 2,
            ):
                if not any(row[i] is not None for i in relevant):
                    continue
                if row_no > MAX_ROWS:
                    raise WorkbookError(f"{spec.code}: limite de {MAX_ROWS} linhas excedido.")
                try:
                    yr = number(row[dims["ano"]])
                    if yr is None or yr != int(yr) or not 1900 <= yr <= 2200:
                        raise WorkbookError("Ano ausente ou inválido.")
                    year = int(yr)
                    month = None
                    if has_month:
                        text = norm(row[dims["mes"]])
                        if text not in MONTHS:
                            raise WorkbookError("Mês ausente ou inválido.")
                        month = MONTHS.index(text) + 1
                    house = str(row[dims["casa"]]).strip() if "casa" in dims and row[dims["casa"]] is not None else None
                    if "casa" in dims and not house:
                        raise WorkbookError("Casa ausente.")
                    location = (
                        str(row[dims["unidade"]]).strip()
                        if "unidade" in dims and row[dims["unidade"]] is not None
                        else None
                    )
                    for col in columns:
                        value = number(row[col])
                        category = None
                        if spec.kind == "wide":
                            category = re.sub(r"^Número de colaboradores\s*", "", str(preview[header][col]), flags=re.I)
                            category = re.sub(r"^(do gênero|de gênero)\s+", "", category, flags=re.I)
                        elif "category" in dims:
                            if row[dims["category"]] is None:
                                raise WorkbookError("Categoria ausente.")
                            category = str(row[dims["category"]]).strip()
                            if spec.code == "A.5.1.1":
                                category = f"Escopo {category}"
                        missing += value is None
                        key = (year, month, house, location, category, value)
                        duplicates[key] += 1
                        records.append(
                            {
                                "year": year,
                                "month": month,
                                "house": house,
                                "location": location,
                                "category": category,
                                "value": value,
                                "cell": f"{get_column_letter(col + 1)}{row_no}",
                            }
                        )
                except WorkbookError as error:
                    raise WorkbookError(f"{spec.code}, linha {row_no}: {error}") from None
            repeated = sum(v - 1 for v in duplicates.values() if v > 1)
            if missing:
                warnings.append(f"{spec.code}: {missing} valores ausentes; não convertidos em zero.")
            if repeated:
                warnings.append(f"{spec.code}: {repeated} registros iguais preservados; revisar na fonte.")
            if not records:
                warnings.append(f"{spec.code}: estrutura válida, sem registros.")
            if has_month:
                periods: dict[tuple[Any, ...], set[int]] = {}
                for record in records:
                    periods.setdefault((record["year"], record["house"]), set()).add(record["month"])
                incomplete = sum(len(m) < 12 for m in periods.values())
                if incomplete:
                    warnings.append(f"{spec.code}: {incomplete} combinações Casa/ano com menos de 12 meses.")
            if spec.kind == "ratio":
                warnings.append(
                    f"{spec.code}: valor anual consolidado, sem recorte por Casa; proporções não são somadas."
                )
            indicators.append(
                {
                    "code": spec.code,
                    "title": spec.title,
                    "unit": display_unit,
                    "sourceUnit": unit,
                    "sheet": sheet.title,
                    "kind": spec.kind,
                    "monthly": has_month,
                    "houseFilter": "casa" in dims,
                    "records": records,
                    "missing": missing,
                    "duplicates": repeated,
                }
            )
        all_records = [r for i in indicators for r in i["records"]]
        return {
            "indicators": indicators,
            "warnings": warnings,
            "summary": {
                "indicators": len(indicators),
                "records": len(all_records),
                "missing": sum(r["value"] is None for r in all_records),
                "years": sorted({r["year"] for r in all_records}),
                "houses": sorted({r["house"] for r in all_records if r["house"]}),
            },
        }
    finally:
        workbook.close()
